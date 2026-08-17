"""Client-facing safety layer for KORGAN Telegram.

Legal verification details remain available to KORGAN's internal logs and gates,
but clients are never asked to choose statutes, interpret citation-audit output,
or act on implementation labels such as NEEDS_VERIFICATION / source-bound / QA.

This module also bridges the legacy final citation audit to the verified SQLite
Adilet corpus.  The old audit predates the SQLite corpus and otherwise reports a
false "provision absent" even after the same provision was loaded from Adilet.
"""

from __future__ import annotations

import io
import logging
import sqlite3
from pathlib import Path
from typing import Any

from aiogram import Bot
from aiogram.types import BufferedInputFile
from docx import Document

from korgan.legal.corpus import (
    ACT_GK_GENERAL,
    ACT_GK_SPECIAL,
    ACT_GPK,
    ACT_TAX_DUTY,
    DEFAULT_DB_PATH,
)
from korgan.provision_corpus import ProvisionRecord, VERIFIED

LOGGER = logging.getLogger(__name__)

_GENERIC_CHECK_MESSAGE = (
    "Документ пока не прошёл автоматическую юридическую проверку. "
    "Повторите подготовку документа. Если проверка снова не завершится, "
    "можно передать дело персональному юристу."
)

_GATE_MARKERS = (
    "как поступить:",
    "не понял, что сделать с замечаниями",
    "уточните, что сделать с ними",
    "что именно нужно сделать?",
    "что именно не прошло проверку:",
    "договор не выпущен: обнаружены дефекты правовых ссылок",
    "в замечаниях не было",
)
_INTERNAL_TOKENS = (
    "NEEDS_VERIFICATION",
    "KORGAN QA STATUS",
    "KORGAN QUALITY",
    "PRELIMINARY",
    "source-bound",
)

_ACT_IDS: dict[str, tuple[str, ...]] = {
    "ГК РК": (ACT_GK_GENERAL, ACT_GK_SPECIAL),
    "ГПК РК": (ACT_GPK,),
    "НК РК": (ACT_TAX_DUTY,),
}


def sanitize_client_text(value: str | None) -> str | None:
    """Remove internal workflow vocabulary from Telegram output.

    Normal legal explanations may still cite statutes.  What is hidden is the
    *workflow asking the client to decide verification mechanics* and internal
    machine labels.
    """
    if value is None:
        return None
    text = str(value)
    lowered = text.lower()

    if any(marker in lowered for marker in _GATE_MARKERS):
        return _GENERIC_CHECK_MESSAGE

    # Document captions used to expose QA/verification state and then dump the
    # internal checklist.  Replace the whole caption with a normal client result.
    upper = text.upper()
    if any(token.upper() in upper for token in ("KORGAN QUALITY", "PRELIMINARY")) or (
        ("✅ VERIFIED" in upper or "⚠️ NEEDS_VERIFICATION" in upper)
        and any(word in lowered for word in ("иск", "договор", "отзыв", "word", ".docx"))
    ):
        if "отзыв" in lowered:
            return "✅ Отзыв на иск сформирован в Word (.docx).\n\nПеред подачей проверьте реквизиты, суммы и приложения."
        if "договор" in lowered:
            return "✅ Проект договора сформирован в Word (.docx).\n\nПеред подписанием проверьте реквизиты, суммы и приложения."
        return "✅ Проект иска сформирован в Word (.docx).\n\nПеред подачей проверьте реквизиты, суммы и приложения."

    # Defence in depth for help text and any future client copy.
    text = text.replace("NEEDS_VERIFICATION", "дополнительной проверкой системы")
    text = text.replace("source-bound", "по официальному источнику")
    text = text.replace("корпус KORGAN", "проверенная правовая база")
    text = text.replace("KORGAN QA STATUS", "")
    return text


def _remove_paragraph(paragraph: Any) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def clean_client_docx(data: bytes) -> bytes:
    """Strip pure internal QA labels from a DOCX before Telegram delivery."""
    try:
        document = Document(io.BytesIO(data))
    except Exception:
        return data

    changed = False
    for paragraph in list(document.paragraphs):
        upper = (paragraph.text or "").strip().upper()
        if upper.startswith("KORGAN QA STATUS:") or upper in {
            "PRELIMINARY DRAFT",
            "LAWYER-REVIEW DRAFT",
            "READY FOR FINAL HUMAN REVIEW",
        }:
            _remove_paragraph(paragraph)
            changed = True

    if not changed:
        return data
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _clean_upload(document: Any) -> Any:
    if not isinstance(document, BufferedInputFile):
        return document
    filename = str(getattr(document, "filename", "") or "")
    payload = getattr(document, "data", None)
    if not filename.lower().endswith(".docx") or not isinstance(payload, (bytes, bytearray)):
        return document
    cleaned = clean_client_docx(bytes(payload))
    return BufferedInputFile(cleaned, filename=filename)


class ClientSafeBot(Bot):
    """Bot transport that prevents internal legal-QA vocabulary from leaking."""

    async def send_message(self, chat_id: Any, text: str, *args: Any, **kwargs: Any) -> Any:
        return await super().send_message(chat_id, sanitize_client_text(text) or "", *args, **kwargs)

    async def send_document(self, chat_id: Any, document: Any, *args: Any, **kwargs: Any) -> Any:
        if "caption" in kwargs:
            kwargs["caption"] = sanitize_client_text(kwargs.get("caption"))
        return await super().send_document(chat_id, _clean_upload(document), *args, **kwargs)

    async def edit_message_text(self, text: str, *args: Any, **kwargs: Any) -> Any:
        return await super().edit_message_text(sanitize_client_text(text) or "", *args, **kwargs)

    async def edit_message_caption(self, *args: Any, **kwargs: Any) -> Any:
        if "caption" in kwargs:
            kwargs["caption"] = sanitize_client_text(kwargs.get("caption"))
        return await super().edit_message_caption(*args, **kwargs)


def _lookup_local(
    act: str,
    article: str,
    part: str = "",
    *,
    db_path: Path | str = DEFAULT_DB_PATH,
) -> ProvisionRecord | None:
    """Resolve a legacy citation-audit reference from the verified SQLite corpus."""
    path = Path(db_path)
    if not path.exists():
        return None
    act_ids = _ACT_IDS.get(act, ())
    if not act_ids:
        return None

    placeholders = ",".join("?" for _ in act_ids)
    params: list[Any] = [*act_ids, str(article)]
    sql = (
        "SELECT p.act_id, p.article_no, p.item_no, p.body, p.edition_date, p.url "
        "FROM provisions p WHERE p.act_id IN (" + placeholders + ") AND p.article_no = ?"
    )
    if part:
        sql += " AND (p.item_no = ? OR p.item_no IS NULL)"
        params.append(str(part))
    sql += " ORDER BY p.sort_key"

    try:
        connection = sqlite3.connect(path)
        connection.row_factory = sqlite3.Row
        rows = connection.execute(sql, params).fetchall()
    except sqlite3.Error:
        LOGGER.exception("CLIENT_SAFE local citation lookup failed act=%s article=%s part=%s", act, article, part)
        return None
    finally:
        try:
            connection.close()  # type: ignore[possibly-undefined]
        except Exception:
            pass

    if not rows:
        return None

    pieces: list[str] = []
    for row in rows:
        body = str(row["body"] or "").strip()
        item = str(row["item_no"] or "").strip()
        if body:
            pieces.append(f"{item}. {body}" if item and not part else body)
    text = "\n".join(pieces).strip()
    if not text:
        return None

    first = rows[0]
    verified_on = max(str(row["edition_date"] or "") for row in rows)
    return ProvisionRecord(
        act=act,
        act_aliases=(act,),
        article=str(article),
        part=str(part or ""),
        text=text,
        source_url=str(first["url"] or ""),
        verified_on=verified_on,
        level=VERIFIED,
        provenance="Загружено KORGAN из официальной русской редакции Adilet в локальный SQLite-корпус.",
    )


def install_client_safe_runtime() -> None:
    """Install production-only UI and citation-audit bridges."""
    from korgan import bot as base_bot
    from korgan import citation_audit, provision_corpus

    if getattr(base_bot, "_client_safe_runtime_installed", False):
        return

    original_static_lookup = provision_corpus.lookup

    def local_aware_lookup(act: str, article: str, part: str = "") -> ProvisionRecord | None:
        local = _lookup_local(act, article, part)
        if local is not None:
            return local
        return original_static_lookup(act, article, part)

    # audit_citations imported lookup directly, therefore patch both bindings.
    provision_corpus.lookup = local_aware_lookup
    citation_audit.lookup = local_aware_lookup

    original_builder = base_bot.build_claim_docx

    def safe_claim_builder(draft: Any) -> bytes:
        return clean_client_docx(original_builder(draft))

    base_bot.build_claim_docx = safe_claim_builder

    async def safe_enter_verification_gate(message: Any, state: Any, draft: Any, report: Any) -> None:
        issues = list(getattr(report, "blocking", []) or [])[:12]
        LOGGER.warning("CLAIM_RELEASE_BLOCKED_INTERNAL issues=%s", issues)
        # Never put a client into a state where they are expected to choose or
        # waive a provision.  Legal citation decisions belong to KORGAN/lawyer.
        await state.update_data(mode="main", gate_issues=[], claim_draft=None)
        await message.answer(_GENERIC_CHECK_MESSAGE, reply_markup=base_bot.MENU)

    async def safe_gate_reply(message: Any, state: Any, data: dict[str, Any]) -> None:
        LOGGER.info("STALE_VERIFICATION_GATE_CLEARED telegram_user_id=%s", getattr(getattr(message, "from_user", None), "id", None))
        await state.update_data(mode="main", gate_issues=[], claim_draft=None)
        await message.answer(_GENERIC_CHECK_MESSAGE, reply_markup=base_bot.MENU)

    base_bot._enter_verification_gate = safe_enter_verification_gate
    base_bot._handle_verification_gate_reply = safe_gate_reply
    base_bot._client_safe_runtime_installed = True
    LOGGER.info("KORGAN client-safe legal UI installed")
