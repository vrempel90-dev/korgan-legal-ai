"""Регрессия фазы 1: числа иска приходят только из детерминированного расчёта.

Два кейса взяты из боевых обращений, на которых расхождение было замечено
вручную: поставка с одной частичной оплатой и поставка с двумя оплатами и
договорным пределом неустойки. Оба проверяются целиком — от разбора материалов
до сумм в готовом Word, — потому что ошибка здесь проявляется не в отдельной
функции, а в расхождении между разделами документа.

Тест намеренно сверяет КАЖДОЕ число отдельно, а не только итог. Итог сходится и
у неверного расчёта: достаточно, чтобы одна ошибка компенсировала другую, и
проверка «цена иска совпала» пропустит и неверную базу, и неверный период.
"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

import pytest
from docx import Document

from korgan.claim_calculation_contract import FieldStatus
from korgan.claim_calculator import build_claim_calculation, try_calculator_authority
from korgan.claim_docx import build_claim_docx
from korgan.claim_financials import CapBase, extract_case_financials
from korgan.claim_money_ledger import build_claim_money_ledger
from korgan.late_interest_hotfix import CALCULATOR_NOTE_PREFIX, _apply_verified_penalty
from korgan.legal_calc import format_kzt
from korgan.legal_types import ClaimDraft, LegalResearch, VerificationStatus

CASE_A_CONTEXT = """Файл: KAZ_INDUSTRY_TRADE_postavka.docx
Истец: ТОО «KAZ INDUSTRY TRADE», БИН 010140001230, г. Алматы, ул. Абая, 150, офис 12
Ответчик: ТОО «АлматыСтройСнаб», БИН 020240005675, г. Алматы, пр. Райымбека, 208
Договор поставки № 14/2026 от 02.02.2026. Стоимость поставки составила 8 750 000 тенге.
Товар поставлен в полном объёме 20.02.2026, накладная № 77 от 20.02.2026.
Срок оплаты по договору — до 10.03.2026 включительно.
07.04.2026 ответчик частично оплатил 2 000 000 тенге в счёт погашения задолженности.
Остаток основного долга составляет 6 750 000 тенге.
Пунктом 5.2 договора предусмотрена неустойка в размере 0,1% от суммы задолженности за каждый день просрочки.
Прошу взыскать основной долг и неустойку. Расчёт неустойки произвести по 01.09.2026 включительно.
"""

CASE_B_CONTEXT = """Файл: ASTANA_SUPPLY_GROUP_postavka.docx
Истец: ТОО «ASTANA SUPPLY GROUP», БИН 030340009019, г. Астана, ул. Кунаева, 12
Ответчик: ТОО «Сарыарка Логистик», БИН 040440003451, г. Астана, ул. Сейфуллина, 3
Договор поставки № 7 от 01.02.2026. Стоимость поставки составила 7 200 000 тенге.
Срок оплаты по договору — до 10.03.2026 включительно.
20.03.2026 ответчик частично оплатил 1 200 000 тенге в счёт погашения задолженности.
15.04.2026 ответчик частично оплатил 1 000 000 тенге в счёт погашения задолженности.
Остаток основного долга составляет 5 000 000 тенге.
Пунктом 6.3 договора предусмотрена неустойка в размере 0,1% от фактической просроченной задолженности за каждый день просрочки, но не более 10% первоначальной стоимости поставленного товара.
Прошу взыскать основной долг и неустойку. Расчёт неустойки произвести по 02.09.2026 включительно.
"""

# Ожидаемые значения посчитаны вручную по материалам дела и зафиксированы здесь
# как контрольные. Они не выводятся из кода продукта: иначе тест сверял бы
# реализацию сама с собой и молчал бы ровно при той ошибке, ради которой написан.
CASE_A_EXPECTED = {
    "contract_value": 8_750_000,
    "principal": 6_750_000,
    # 11.03.2026—06.04.2026 — 27 дней от 8 750 000; 07.04.2026—01.09.2026 — 148 дней от 6 750 000.
    "penalty_intervals": (
        (date(2026, 3, 11), date(2026, 4, 6), 27, 8_750_000, 236_250),
        (date(2026, 4, 7), date(2026, 9, 1), 148, 6_750_000, 999_000),
    ),
    "penalty": 1_235_250,
    "claim_price": 7_985_250,
    "state_duty": 239_558,   # 3% от 7 985 250 для юридического лица
    "total_claim": 8_224_808,
    "capped": False,
}

CASE_B_EXPECTED = {
    "contract_value": 7_200_000,
    "principal": 5_000_000,
    "penalty_intervals": (
        (date(2026, 3, 11), date(2026, 3, 19), 9, 7_200_000, 64_800),
        (date(2026, 3, 20), date(2026, 4, 14), 26, 6_000_000, 156_000),
        (date(2026, 4, 15), date(2026, 9, 2), 141, 5_000_000, 705_000),
    ),
    "penalty_raw": 925_800,
    "cap": 720_000,          # 10% первоначальной стоимости поставки
    "penalty": 720_000,
    "claim_price": 5_720_000,
    "state_duty": 171_600,   # 3% от 5 720 000
    "total_claim": 5_891_600,
    "capped": True,
}

CASES = (
    pytest.param(CASE_A_CONTEXT, CASE_A_EXPECTED, date(2026, 9, 1), id="kaz-industry-trade"),
    pytest.param(CASE_B_CONTEXT, CASE_B_EXPECTED, date(2026, 9, 2), id="astana-supply-group"),
)


def _research() -> LegalResearch:
    return LegalResearch(
        status=VerificationStatus.VERIFIED,
        applicable_law=[],
        procedural_requirements=[],
        verified_claims=[],
        unverified_claims=[],
        source_urls=[],
        notes=[],
    )


def _model_draft(**overrides) -> ClaimDraft:
    """Черновик в том виде, в каком его отдаёт модель — со своими числами.

    Суммы здесь намеренно неверные. Тест проверяет не то, что модель ошиблась,
    а то, что её число не доживает до документа ни в одном разделе.
    """
    base = dict(
        status=VerificationStatus.NEEDS_VERIFICATION,
        title="ИСКОВОЕ ЗАЯВЛЕНИЕ о взыскании задолженности по договору поставки",
        court="Специализированный межрайонный экономический суд",
        claimant=["ТОО «KAZ INDUSTRY TRADE», БИН 010140001230"],
        defendant=["ТОО «АлматыСтройСнаб», БИН 020240005675"],
        price_of_claim="9 999 999 тенге",
        facts=["Ответчик не оплатил поставленный товар."],
        legal_basis=["Согласно договору поставки ответчик обязан оплатить товар."],
        requests=[
            "Взыскать с ответчика основной долг в размере 8 750 000 тенге.",
            "Взыскать с ответчика неустойку в размере 1 500 000 тенге.",
            "Взыскать с ответчика расходы по уплате государственной пошлины.",
        ],
        attachments=["Копия договора поставки № 14/2026 от 02.02.2026"],
        verification_notes=[],
        source_urls=[],
    )
    base.update(overrides)
    return ClaimDraft(**base)


def _drafted(context: str, filing_date: date) -> ClaimDraft:
    draft = _model_draft()
    _apply_verified_penalty(context, _research(), draft, filing_date=filing_date)
    return draft


def _docx_text(draft: ClaimDraft) -> str:
    document = Document(io.BytesIO(build_claim_docx(draft)))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


# --------------------------------------------------------------------------
# Разбор материалов дела
# --------------------------------------------------------------------------


@pytest.mark.parametrize(("context", "expected", "filing_date"), CASES)
def test_inputs_are_read_from_case_materials(context, expected, filing_date) -> None:
    financials = extract_case_financials(context)

    assert financials.contract_value == expected["contract_value"]
    assert financials.principal == expected["principal"]
    assert financials.due_date == date(2026, 3, 10)
    assert financials.calculation_end == filing_date
    assert financials.penalty_rate_per_day == Decimal("0.1")
    assert financials.missing == ()


def test_case_a_reads_the_single_partial_payment() -> None:
    financials = extract_case_financials(CASE_A_CONTEXT)

    assert [(event.on, -int(event.delta)) for event in financials.payments] == [
        (date(2026, 4, 7), 2_000_000)
    ]
    assert financials.cap_percent is None


def test_case_b_reads_both_payments_and_the_contractual_cap() -> None:
    financials = extract_case_financials(CASE_B_CONTEXT)

    assert [(event.on, -int(event.delta)) for event in financials.payments] == [
        (date(2026, 3, 20), 1_200_000),
        (date(2026, 4, 15), 1_000_000),
    ]
    # Предел считается от первоначальной стоимости товара, а не от остатка долга.
    assert financials.cap_percent == Decimal("10")
    assert financials.cap_base is CapBase.CONTRACT_VALUE
    assert financials.cap_amount == CASE_B_EXPECTED["cap"]


def test_stated_principal_that_contradicts_the_payments_is_refused() -> None:
    """Названный остаток и арифметика платежей проверяют друг друга."""
    broken = CASE_A_CONTEXT.replace(
        "Остаток основного долга составляет 6 750 000 тенге.",
        "Остаток основного долга составляет 6 000 000 тенге.",
    )

    financials = extract_case_financials(broken)

    assert financials.principal is None
    assert financials.principal_conflict is True
    assert any("не сходится" in note for note in financials.missing)


# --------------------------------------------------------------------------
# Структурированный расчёт
# --------------------------------------------------------------------------


@pytest.mark.parametrize(("context", "expected", "filing_date"), CASES)
def test_every_calculated_field_matches_the_manual_result(context, expected, filing_date) -> None:
    outcome = build_claim_calculation(
        context,
        extract_case_financials(context),
        filing_date=filing_date,
        penalty_claimed=True,
    )
    calculation = outcome.calculation

    assert calculation.ready is True
    assert calculation.principal.value == expected["principal"]
    assert calculation.penalty.value == expected["penalty"]
    assert calculation.claim_price.value == expected["claim_price"]
    assert calculation.state_duty.value == expected["state_duty"]
    assert calculation.total_claim.value == expected["total_claim"]
    assert calculation.insufficient_fields() == ()


@pytest.mark.parametrize(("context", "expected", "filing_date"), CASES)
def test_penalty_periods_days_and_bases_match(context, expected, filing_date) -> None:
    """Периоды и дни проверяются отдельно от итога.

    Неверный период с компенсирующей ошибкой в базе даёт правильную сумму и
    неправильный расчёт. Суд читает таблицу, а не только итог.
    """
    outcome = build_claim_calculation(
        context,
        extract_case_financials(context),
        filing_date=filing_date,
        penalty_claimed=True,
    )
    assert outcome.penalty is not None

    actual = tuple(
        (i.period_from, i.period_to, i.days, i.principal, i.subtotal)
        for i in outcome.penalty.intervals
    )
    assert actual == expected["penalty_intervals"]


@pytest.mark.parametrize(("context", "expected", "filing_date"), CASES)
def test_claim_price_equals_the_sum_of_its_components(context, expected, filing_date) -> None:
    outcome = build_claim_calculation(
        context,
        extract_case_financials(context),
        filing_date=filing_date,
        penalty_claimed=True,
    )
    calculation = outcome.calculation

    assert calculation.claim_price.value == (
        int(calculation.principal.value or 0) + int(calculation.penalty.value or 0)
    )
    assert calculation.total_claim.value == (
        int(calculation.claim_price.value or 0) + int(calculation.state_duty.value or 0)
    )


def test_contractual_cap_limits_the_penalty_in_case_b() -> None:
    outcome = build_claim_calculation(
        CASE_B_CONTEXT,
        extract_case_financials(CASE_B_CONTEXT),
        filing_date=date(2026, 9, 2),
        penalty_claimed=True,
    )
    assert outcome.penalty is not None

    assert outcome.penalty.raw_total == CASE_B_EXPECTED["penalty_raw"]
    assert outcome.penalty.capped is True
    assert outcome.penalty.cap_amount == CASE_B_EXPECTED["cap"]
    assert outcome.calculation.penalty.value == CASE_B_EXPECTED["cap"]


def test_case_a_penalty_is_not_capped() -> None:
    outcome = build_claim_calculation(
        CASE_A_CONTEXT,
        extract_case_financials(CASE_A_CONTEXT),
        filing_date=date(2026, 9, 1),
        penalty_claimed=True,
    )
    assert outcome.penalty is not None
    assert outcome.penalty.capped is False


# --------------------------------------------------------------------------
# Числа модели не доживают до документа
# --------------------------------------------------------------------------


@pytest.mark.parametrize(("context", "expected", "filing_date"), CASES)
def test_model_written_amounts_are_replaced_everywhere(context, expected, filing_date) -> None:
    draft = _drafted(context, filing_date)
    text = _docx_text(draft)

    for invented in ("9 999 999", "1 500 000"):
        assert invented not in text, f"число модели {invented} осталось в документе"
    assert format_kzt(expected["claim_price"]) in text
    assert format_kzt(expected["principal"]) in text
    assert format_kzt(expected["penalty"]) in text


@pytest.mark.parametrize(("context", "expected", "filing_date"), CASES)
def test_the_same_amount_appears_in_every_section(context, expected, filing_date) -> None:
    """Шапка, раздел расчёта и просительная часть говорят одно число."""
    draft = _drafted(context, filing_date)

    assert draft.price_of_claim == format_kzt(expected["claim_price"])
    assert draft.state_duty.startswith(format_kzt(expected["state_duty"]))

    calculation_text = "\n".join(draft.calculation)
    assert f"Итого цена иска: {format_kzt(expected['claim_price'])}" in calculation_text
    assert f"Итого ко взысканию с ответчика: {format_kzt(expected['total_claim'])}" in calculation_text

    prayer = "\n".join(draft.requests)
    assert format_kzt(expected["principal"]) in prayer
    assert format_kzt(expected["penalty"]) in prayer


@pytest.mark.parametrize(("context", "expected", "filing_date"), CASES)
def test_money_ledger_reads_back_the_same_claim_price(context, expected, filing_date) -> None:
    """Второй, независимый обход просительной части даёт ту же цену иска.

    Финалайзер пересчитывает цену иска своим разбором. Если он прочитает из
    текста другое число, документ снова начнёт утверждать две суммы — уже без
    участия модели.
    """
    draft = _drafted(context, filing_date)

    ledger = build_claim_money_ledger(list(draft.requests))

    assert ledger.unresolved_requests == []
    assert ledger.total == expected["claim_price"]


@pytest.mark.parametrize(("context", "expected", "filing_date"), CASES)
def test_structured_result_is_attached_to_the_draft(context, expected, filing_date) -> None:
    draft = _drafted(context, filing_date)

    result = draft.calculation_result
    assert result["ready"] is True
    assert result["principal"]["value"] == expected["principal"]
    assert result["penalty"]["value"] == expected["penalty"]
    assert result["claim_price"]["value"] == expected["claim_price"]
    assert result["state_duty"]["value"] == expected["state_duty"]
    assert result["total_claim"]["value"] == expected["total_claim"]
    # Основание каждого числа сохранено: юрист сверяет его с материалами дела.
    assert result["principal"]["source"]
    assert result["inputs"]["due_date"] == "2026-03-10"


@pytest.mark.parametrize(("context", "expected", "filing_date"), CASES)
def test_no_verification_markers_are_left_in_a_complete_case(context, expected, filing_date) -> None:
    draft = _drafted(context, filing_date)
    text = _docx_text(draft)

    assert "[ТРЕБУЕТ РАСЧ" not in text.upper()
    assert "{{" not in text
    assert not [note for note in draft.verification_notes if note.startswith(CALCULATOR_NOTE_PREFIX)]


# --------------------------------------------------------------------------
# Дефицит данных — состояние, а не текст
# --------------------------------------------------------------------------


def test_missing_due_date_leaves_the_penalty_unfilled_not_guessed() -> None:
    context = CASE_A_CONTEXT.replace("Срок оплаты по договору — до 10.03.2026 включительно.\n", "")

    outcome = build_claim_calculation(
        context, extract_case_financials(context), filing_date=date(2026, 9, 1), penalty_claimed=True
    )

    assert outcome.calculation.penalty.status is FieldStatus.INSUFFICIENT_DATA
    assert outcome.calculation.penalty.value is None
    assert any("срок оплаты" in reason for reason in outcome.calculation.penalty.missing)
    assert outcome.calculation.ready is False


def test_unreadable_partial_payment_stops_the_calculation() -> None:
    context = CASE_A_CONTEXT.replace(
        "07.04.2026 ответчик частично оплатил 2 000 000 тенге в счёт погашения задолженности.",
        "Ответчик частично оплатил часть долга, точная дата и сумма платежа не установлены.",
    )

    financials = extract_case_financials(context)
    outcome = build_claim_calculation(
        context, financials, filing_date=date(2026, 9, 1), penalty_claimed=True
    )

    assert financials.payments_unclear is True
    assert outcome.calculation.penalty.status is FieldStatus.INSUFFICIENT_DATA


def test_insufficient_data_never_authors_the_document() -> None:
    """При дефиците данных калькулятор не берёт авторство и не пишет чисел."""
    context = CASE_A_CONTEXT.replace("Срок оплаты по договору — до 10.03.2026 включительно.\n", "")
    draft = _model_draft()

    assert (
        try_calculator_authority(
            context, draft, filing_date=date(2026, 9, 1), penalty_claimed=True
        )
        is None
    )


def test_a_claim_with_other_money_relief_is_left_to_the_existing_path() -> None:
    """Убытки калькулятор не считает — значит и остальные числа не его.

    Посчитать долг и неустойку, а убытки оставить из текста модели — это цена
    иска, собранная из двух источников. Она сойдётся только случайно.
    """
    draft = _model_draft(
        requests=[
            "Взыскать с ответчика основной долг в размере 8 750 000 тенге.",
            "Взыскать с ответчика убытки в размере 400 000 тенге.",
        ]
    )

    assert (
        try_calculator_authority(
            CASE_A_CONTEXT, draft, filing_date=date(2026, 9, 1), penalty_claimed=True
        )
        is None
    )


def test_lawyer_message_stays_out_of_the_court_text() -> None:
    context = CASE_A_CONTEXT.replace("Срок оплаты по договору — до 10.03.2026 включительно.\n", "")
    draft = _model_draft()

    _apply_verified_penalty(context, _research(), draft, filing_date=date(2026, 9, 1))

    body = "\n".join([draft.title, draft.price_of_claim, *draft.facts, *draft.legal_basis, *draft.requests])
    assert CALCULATOR_NOTE_PREFIX not in body


# --------------------------------------------------------------------------
# Плейсхолдеры модели
# --------------------------------------------------------------------------


def test_placeholders_are_substituted_with_calculated_values() -> None:
    draft = _model_draft(
        price_of_claim="{{claim_price}}",
        requests=[
            "Взыскать с ответчика основной долг в размере {{principal_amount}}.",
            "Взыскать с ответчика неустойку в размере {{penalty_amount}}.",
        ],
        facts=["Задолженность ответчика составляет {{principal_amount}}."],
    )

    outcome = try_calculator_authority(
        CASE_A_CONTEXT, draft, filing_date=date(2026, 9, 1), penalty_claimed=True
    )

    assert outcome is not None
    assert "{{" not in "\n".join([draft.price_of_claim, *draft.facts, *draft.requests])
    assert format_kzt(CASE_A_EXPECTED["principal"]) in draft.facts[0]


def test_penalty_the_client_never_asked_for_leaves_no_trace() -> None:
    """Расчёт не наследует позицию модели.

    Модель написала иск о долге и неустойке, клиент просил только долг.
    Калькулятор считает заявленное, а упоминание неустойки уходит из всех
    разделов — иначе документ просит одно, а рассказывает о другом.
    """
    context = (
        "Истец: ТОО «KAZ INDUSTRY TRADE», БИН 010140001230. "
        "Ответчик: ТОО «АлматыСтройСнаб», БИН 020240005675. "
        "Прошу взыскать только основной долг 6 750 000 тенге."
    )
    draft = _model_draft(
        facts=["Задолженность 6 750 000 тенге основного долга и 1 500 000 тенге неустойки."],
        attachments=["Расчёт неустойки на 1 500 000 тенге."],
    )

    _apply_verified_penalty(context, _research(), draft, filing_date=date(2026, 9, 1))

    assert not any("неустой" in item.lower() for item in draft.requests)
    assert not any("неустой" in item.lower() for item in draft.facts)
    assert not any("неустой" in item.lower() for item in draft.attachments)
    assert draft.late_interest == ""
    assert draft.price_of_claim == format_kzt(6_750_000)
    assert draft.calculation_result["penalty"]["status"] == "not_applicable"
    assert draft.calculation_result["claim_price"]["value"] == 6_750_000


def test_unresolved_placeholder_never_reaches_the_document() -> None:
    """Токен без подстановки снимается вместе со своей строкой.

    Так бывает, когда расчёт этого поля не состоялся. Напечатать
    «{{penalty_amount}}» в иске нельзя, а подставить вместо него число модели —
    тем более: ровно от этого числа калькулятор и защищает.
    """
    context = CASE_A_CONTEXT.replace("Срок оплаты по договору — до 10.03.2026 включительно.\n", "")
    draft = _model_draft(
        price_of_claim="{{claim_price}}",
        requests=[
            "Взыскать с ответчика основной долг в размере {{principal_amount}}.",
            "Взыскать с ответчика неустойку в размере {{penalty_amount}}.",
            "Взыскать с ответчика расходы по уплате государственной пошлины.",
        ],
        facts=["Задолженность ответчика составляет {{principal_amount}}."],
    )

    _apply_verified_penalty(context, _research(), draft, filing_date=date(2026, 9, 1))

    rendered = _docx_text(draft)
    assert "{{" not in rendered
    # Поле цены иска не несёт токена. Продуктовый маркер незавершённого расчёта
    # на его месте — существующее поведение выпуска: он переводит документ в
    # предварительный статус, и его снятие относится к hard-gate фазы 3.
    assert "{{" not in draft.price_of_claim
    assert not any("{{" in item for item in draft.requests + draft.facts)
    assert draft.status is VerificationStatus.NEEDS_VERIFICATION
    assert any(note.startswith(CALCULATOR_NOTE_PREFIX) for note in draft.verification_notes)


def test_title_keeps_its_meaning_when_a_placeholder_is_stripped() -> None:
    """Иск без цифры в заголовке остаётся нормальным иском."""
    context = CASE_A_CONTEXT.replace("Срок оплаты по договору — до 10.03.2026 включительно.\n", "")
    draft = _model_draft(title="ИСКОВОЕ ЗАЯВЛЕНИЕ о взыскании {{claim_price}}")

    _apply_verified_penalty(context, _research(), draft, filing_date=date(2026, 9, 1))

    assert draft.title == "ИСКОВОЕ ЗАЯВЛЕНИЕ о взыскании"
    assert "{{" not in draft.title
