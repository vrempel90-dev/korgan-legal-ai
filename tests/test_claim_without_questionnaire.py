"""Приёмка: одно сообщение клиента → один документ, с пометками вместо анкеты.

Тестовое дело — реальное: клиент перевёл 2 300 000 тенге предоплаты за ремонт,
работы не выполнены, деньги не возвращены. Раньше такое сообщение начинало
диалог «не хватает: дата рождения истца… банковские реквизиты…», и черновик
клиент видел через три-четыре раунда. Здесь проверяется, что документ приходит
сразу, а недостающие формальные реквизиты стоят в нём как
`[ТРЕБУЕТ УТОЧНЕНИЯ: ...]`.

Второй сюжет — правовое обоснование. На этом же деле в документ попадала норма
о приёмке результата работ: она про обязанность заказчика, а не про основание
взыскания денег с подрядчика. Проверяется, что такая подмена ловится, а
согласованная пользователем статья не исчезает при пересборке документа.

Сеть не используется: подменяется `korgan.bot.service`.

Сюжет со шлюзом проверки (`verification_gate`) — это собственный диалог
`korgan/bot.py`. В боевом рантайме его намеренно заменяет
`client_safe_ui.install_client_safe_runtime()`; боевой контракт закреплён в
`tests/test_client_safe_gate_supersedes_waiver_flow.py`. Фикстура `bot_own_gate`
закрепляет предмет проверки за реализацией bot.py, чтобы результат не зависел
от того, импортировал ли другой тестовый модуль `korgan.strict_bot` раньше.
"""

from __future__ import annotations

import asyncio
import io
from dataclasses import dataclass, field
from types import SimpleNamespace

import pytest
from aiogram.fsm.context import FSMContext
from aiogram.fsm.storage.base import StorageKey
from aiogram.fsm.storage.memory import MemoryStorage
from docx import Document

from korgan import bot as korgan_bot
from korgan.legal_basis_fit import (
    LAWYER_PICK_MARKER,
    NOTE_PREFIX,
    categorize_provision,
    detect_relief,
    legal_basis_defects,
)
from korgan.legal_types import ClaimDraft, LegalResearch, VerificationStatus


@pytest.fixture(autouse=True)
def bot_own_gate(monkeypatch: pytest.MonkeyPatch):
    """Проверять реализацию bot.py, даже если client-safe слой уже установлен."""
    from korgan import client_safe_ui

    original_enter = client_safe_ui._original_enter_verification_gate
    original_reply = client_safe_ui._original_handle_verification_gate_reply
    if original_enter is not None:
        monkeypatch.setattr(korgan_bot, "_enter_verification_gate", original_enter)
    if original_reply is not None:
        monkeypatch.setattr(korgan_bot, "_handle_verification_gate_reply", original_reply)
    yield

# Ровно одно сообщение клиента — так, как он его пишет.
CASE = (
    "Я, Ахметов Руслан Маратович, нанял подрядчика на ремонт квартиры. "
    "Подрядчик — Садыков Тимур Ерланович. "
    "12.02.2026 я перевёл ему предоплату 2 300 000 тенге. "
    "Работы он так и не начал, сроки сорвал, на связь не выходит, деньги не вернул. "
    "Хочу взыскать предоплату обратно."
)

RESEARCH = LegalResearch(
    status=VerificationStatus.NEEDS_VERIFICATION,
    applicable_law=["ГК РК (Особенная часть), глава 32 «Подряд»"],
    procedural_requirements=[],
    verified_claims=[],
    unverified_claims=[],
    source_urls=[],
    notes=[],
)

# Правильное обоснование: отказ от договора и возврат аванса.
BASIS_SUPPORTS_RELIEF = [
    "Подрядчик к выполнению работ не приступил и нарушил согласованные сроки, "
    "что является существенным нарушением договора подряда.",
    "Заказчик вправе отказаться от исполнения договора и потребовать возврата "
    "уплаченного аванса; удержание предоплаты после отказа от договора образует "
    "неосновательное обогащение подрядчика.",
]

# Дефект из отчёта: норма о приёмке как единственное обоснование.
BASIS_ABOUT_ACCEPTANCE = [
    "Заказчик обязан в сроки и в порядке, предусмотренные договором подряда, "
    "с участием подрядчика осмотреть и принять результат выполненной работы.",
]


def _draft(legal_basis: list[str]) -> ClaimDraft:
    return ClaimDraft(
        status=VerificationStatus.NEEDS_VERIFICATION,
        title="ИСКОВОЕ ЗАЯВЛЕНИЕ о взыскании суммы предоплаты",
        court="",
        claimant=["Ахметов Руслан Маратович"],
        defendant=["Садыков Тимур Ерланович"],
        price_of_claim="2 300 000 тенге",
        facts=[
            "12.02.2026 истец перечислил ответчику предоплату в размере 2 300 000 тенге "
            "за выполнение ремонтных работ.",
            "Ответчик к выполнению работ не приступил, результат не передал, "
            "предоплату не возвратил.",
        ],
        legal_basis=list(legal_basis),
        requests=["Взыскать с ответчика в пользу истца 2 300 000 тенге предоплаты."],
        attachments=["Копия платёжного документа от 12.02.2026"],
        verification_notes=[],
        source_urls=[],
        state_duty="23 000 тенге",
    )


@dataclass
class FakeMessage:
    """Minimal Message: records replies and delivered documents."""

    text: str
    sent: list[str] = field(default_factory=list)
    documents: list[bytes] = field(default_factory=list)
    captions: list[str] = field(default_factory=list)

    def __post_init__(self) -> None:
        self.chat = SimpleNamespace(id=1)
        self.from_user = SimpleNamespace(id=1, is_bot=False)
        self.bot = SimpleNamespace(send_chat_action=self._noop)

    async def _noop(self, *args, **kwargs) -> None:
        return None

    async def answer(self, text: str, **kwargs) -> None:
        self.sent.append(text)

    async def answer_document(self, document, **kwargs) -> None:
        self.documents.append(document.data)
        self.captions.append(str(kwargs.get("caption", "")))


class FakeService:
    """Рантайм без сети: черновик задаётся сценарием теста."""

    def __init__(self, drafts: list[ClaimDraft]) -> None:
        self._drafts = list(drafts)
        self.draft_calls = 0

    async def research_case(self, case_context: str, language: str = "ru") -> LegalResearch:
        return RESEARCH

    async def draft_claim(self, case_context, research, language: str = "ru") -> ClaimDraft:
        self.draft_calls += 1
        draft = self._drafts.pop(0) if len(self._drafts) > 1 else self._drafts[0]
        # Каждый проход отдаёт свежий объект, как настоящая модель.
        return _draft(draft.legal_basis) if isinstance(draft, ClaimDraft) else draft


class Dialog:
    """Гоняет реальные хэндлеры korgan.bot поверх настоящего FSM-хранилища."""

    def __init__(self) -> None:
        self.state = FSMContext(
            storage=MemoryStorage(), key=StorageKey(bot_id=1, chat_id=1, user_id=1)
        )
        self.documents: list[bytes] = []
        self.captions: list[str] = []

    async def start_case(self, case: str = CASE) -> None:
        await self.state.set_data(
            {"language": "ru", "documents": [], "facts": [case], "consulted_articles": [], "mode": "main"}
        )

    async def say(self, text: str) -> list[str]:
        message = FakeMessage(text)
        data = await self.state.get_data()
        if data.get("mode") == "verification_gate":
            await korgan_bot._handle_verification_gate_reply(message, self.state, data)
        elif data.get("mode") == "claim_details":
            await korgan_bot._handle_missing_field_answer(message, self.state, data)
        else:
            await korgan_bot.claim_handler(message, self.state)
        self.documents.extend(message.documents)
        self.captions.extend(message.captions)
        return message.sent

    async def mode(self) -> str:
        return str((await self.state.get_data()).get("mode", ""))


def docx_text(payload: bytes) -> str:
    return "\n".join(paragraph.text for paragraph in Document(io.BytesIO(payload)).paragraphs)


def run(coro):
    return asyncio.run(coro)


@pytest.fixture
def service(monkeypatch: pytest.MonkeyPatch):
    def install(*drafts: ClaimDraft) -> FakeService:
        fake = FakeService(list(drafts))
        monkeypatch.setattr(korgan_bot, "service", fake)
        return fake

    return install


# ─── Приёмка 1: один документ вместо анкеты ───────────────────────────────────


def test_one_message_yields_a_document_without_any_rounds(service) -> None:
    """Клиент описал дело один раз — документ приходит сразу."""

    async def scenario() -> None:
        fake = service(_draft(BASIS_SUPPORTS_RELIEF))
        dialog = Dialog()
        await dialog.start_case()

        replies = await dialog.say("подготовь иск")

        assert dialog.documents, "документ не выдан"
        assert fake.draft_calls == 1
        assert await dialog.mode() == "main"
        # Ни одного раунда «пришлите недостающее».
        assert not any("не хватает" in reply for reply in replies)
        assert not any("Пришлите недостающие" in reply for reply in replies)

    run(scenario())


def test_filing_requisites_arrive_as_placeholders_inside_the_document(service) -> None:
    """Дата рождения, ИИН и адреса не спрашиваются — они помечены в файле."""

    async def scenario() -> None:
        service(_draft(BASIS_SUPPORTS_RELIEF))
        dialog = Dialog()
        await dialog.start_case()

        await dialog.say("подготовь иск")

        text = docx_text(dialog.documents[0])
        assert "[ТРЕБУЕТ УТОЧНЕНИЯ: дата рождения истца]" in text
        assert "[ТРЕБУЕТ УТОЧНЕНИЯ: ИИН истца]" in text
        assert "[ТРЕБУЕТ УТОЧНЕНИЯ: адрес места жительства истца]" in text
        assert "[ТРЕБУЕТ УТОЧНЕНИЯ: адрес места жительства ответчика]" in text
        # Суд из материалов не следует — он тоже пометка, а не догадка.
        assert "[ТРЕБУЕТ УТОЧНЕНИЯ: точное наименование суда]" in text
        # И один общий совет: заполнить всё разом в файле.
        assert any("Заполните в файле пометки" in caption for caption in dialog.captions)

    run(scenario())


def test_a_nameless_case_asks_once_and_then_drafts(service) -> None:
    """Критичный пробел спрашивается одним сообщением, ответ сразу даёт документ."""

    async def scenario() -> None:
        service(_draft(BASIS_SUPPORTS_RELIEF))
        dialog = Dialog()
        await dialog.start_case(
            "Перевёл предоплату 2 300 000 тенге за ремонт, работы не выполнены, деньги не вернули."
        )

        question = await dialog.say("подготовь иск")
        assert len(question) == 1
        assert "кто истец" in question[0] and "кто ответчик" in question[0]

        await dialog.say("Ахметов Руслан Маратович, подрядчик Садыков Тимур Ерланович")

        assert dialog.documents, "после единственного ответа документ так и не пришёл"

    run(scenario())


# ─── Приёмка 2: норма должна доказывать заявленное требование ─────────────────


def test_acceptance_of_works_is_not_a_basis_for_recovering_a_prepayment() -> None:
    defects = legal_basis_defects(
        requests=["Взыскать с ответчика 2 300 000 тенге предоплаты."],
        legal_basis=BASIS_ABOUT_ACCEPTANCE,
        context_lines=["предоплата за ремонтные работы не возвращена"],
    )

    assert defects
    assert "приёмке результата работ" in defects[0]
    assert "отказаться от договора" in defects[0]


def test_withdrawal_and_unjust_enrichment_are_accepted_as_the_basis() -> None:
    assert (
        legal_basis_defects(
            requests=["Взыскать с ответчика 2 300 000 тенге предоплаты."],
            legal_basis=BASIS_SUPPORTS_RELIEF,
            context_lines=["предоплата за ремонтные работы не возвращена"],
        )
        == []
    )


@pytest.mark.parametrize(
    "prayer",
    [
        "Взыскать с ответчика 2 300 000 тенге предоплаты.",
        "Взыскать с ответчика 2 300 000 тенге предварительной оплаты.",
        "Взыскать с ответчика сумму аванса 2 300 000 тенге.",
    ],
)
def test_a_prepayment_is_recognised_under_every_wording(prayer: str) -> None:
    """«Предварительная оплата» — та же предоплата: иначе проверка выключалась."""
    relief = detect_relief([prayer], ["оплата за ремонтные работы не возвращена"])

    assert relief is not None and relief.code == "prepayment_refund_works"


def test_relief_and_categories_are_read_from_the_prayer_and_the_text() -> None:
    relief = detect_relief(
        ["Взыскать с ответчика 2 300 000 тенге предоплаты."],
        ["предоплата за ремонт квартиры"],
    )
    assert relief is not None and relief.code == "prepayment_refund_works"
    assert "work_acceptance" in categorize_provision(BASIS_ABOUT_ACCEPTANCE[0])
    assert "contract_withdrawal" in categorize_provision(BASIS_SUPPORTS_RELIEF[1])


def test_document_says_a_lawyer_must_pick_the_provision(service) -> None:
    """Вместо формально похожей нормы — явная пометка о подборе нормы юристом."""

    async def scenario() -> None:
        service(_draft(BASIS_ABOUT_ACCEPTANCE))
        dialog = Dialog()
        await dialog.start_case()

        await dialog.say("подготовь иск")

        text = docx_text(dialog.documents[0])
        assert LAWYER_PICK_MARKER in text
        assert any(NOTE_PREFIX in caption for caption in dialog.captions)

    run(scenario())


def test_a_supporting_provision_is_left_alone(service) -> None:
    async def scenario() -> None:
        service(_draft(BASIS_SUPPORTS_RELIEF))
        dialog = Dialog()
        await dialog.start_case()

        await dialog.say("подготовь иск")

        text = docx_text(dialog.documents[0])
        assert LAWYER_PICK_MARKER not in text
        assert "отказаться от исполнения договора" in text

    run(scenario())


# ─── Приёмка 3: согласованная статья не исчезает ──────────────────────────────

# Норма, которой нет в проверенном корпусе: release gate заблокирует документ,
# и пользователю будет предложено согласиться на пометку.
BASIS_WITH_ARTICLE_616 = [
    "Согласно статье 616 ГК РК заказчик вправе отказаться от исполнения договора "
    "подряда и потребовать возврата уплаченного аванса при существенном нарушении "
    "подрядчиком сроков выполнения работ.",
]


def test_accepted_article_stays_in_the_document_as_needs_verification(service) -> None:
    """Пользователь согласился на пометку — статья остаётся, содержание не утверждается."""

    async def scenario() -> None:
        service(_draft(BASIS_WITH_ARTICLE_616))
        dialog = Dialog()
        await dialog.start_case()

        await dialog.say("подготовь иск")
        assert await dialog.mode() == "verification_gate"
        assert not dialog.documents

        await dialog.say("пометь статью 616 ГК РК как NEEDS_VERIFICATION и продолжи")

        assert dialog.documents, "после согласия документ так и не вышел"
        text = docx_text(dialog.documents[0])
        assert "616" in text
        assert "[ТРЕБУЕТ ПРОВЕРКИ" in text
        # Согласие записано в дело, а не только в текущий черновик.
        stored = (await dialog.state.get_data()).get("accepted_provisions")
        assert stored == [{"act": "ГК РК", "article": "616", "part": ""}]

    run(scenario())


def test_the_accepted_article_survives_a_rebuild_of_the_document(service) -> None:
    """Ключевой дефект: при пересборке статья молча заменялась на более слабую."""

    async def scenario() -> None:
        # Второй проход модели «забывает» статью 616 и уходит в норму о приёмке.
        service(_draft(BASIS_WITH_ARTICLE_616), _draft(BASIS_ABOUT_ACCEPTANCE))
        dialog = Dialog()
        await dialog.start_case()

        await dialog.say("подготовь иск")
        await dialog.say("пометь статью 616 ГК РК как NEEDS_VERIFICATION и продолжи")
        assert dialog.documents

        # Клиент присылает ещё факты — документ собирается заново.
        await dialog.say("подготовь иск")

        text = docx_text(dialog.documents[-1])
        assert "616" in text, "согласованная статья исчезла при пересборке"
        assert "содержание нормы не утверждается" in text
        # Подменившая её норма о приёмке не выдаётся за обоснование требования.
        assert LAWYER_PICK_MARKER in text

    run(scenario())


def test_the_restored_paragraph_reads_as_russian(service) -> None:
    """Пометка попадает в судебный документ — она должна быть согласована по падежам."""

    async def scenario() -> None:
        service(_draft(BASIS_WITH_ARTICLE_616), _draft(BASIS_ABOUT_ACCEPTANCE))
        dialog = Dialog()
        await dialog.start_case()

        await dialog.say("подготовь иск")
        await dialog.say("пометь статью 616 ГК РК как NEEDS_VERIFICATION и продолжи")
        await dialog.say("подготовь иск")

        text = docx_text(dialog.documents[-1])
        assert "содержание статьи 616 ГК РК подлежат сверке" in text
        assert "содержание статья" not in text

    run(scenario())


def test_dropping_an_article_revokes_the_earlier_acceptance(service) -> None:
    """«Убери статью» отменяет прежнее согласие, иначе она вернулась бы обратно."""

    async def scenario() -> None:
        service(_draft(BASIS_WITH_ARTICLE_616))
        dialog = Dialog()
        await dialog.start_case()

        await dialog.say("подготовь иск")
        await dialog.say("пометь статью 616 ГК РК как NEEDS_VERIFICATION и продолжи")
        assert (await dialog.state.get_data()).get("accepted_provisions")

        # Пользователь передумал и вернулся к тому же спорному черновику.
        await dialog.state.update_data(
            mode="verification_gate",
            gate_issues=["статья 616 ГК РК: содержание не подтверждено"],
            claim_draft=korgan_bot._draft_to_state(_draft(BASIS_WITH_ARTICLE_616)),
        )
        await dialog.say("убери статью 616 ГК РК")

        assert (await dialog.state.get_data()).get("accepted_provisions") == []
        assert "616" not in docx_text(dialog.documents[-1])

    run(scenario())
